from io import BytesIO, StringIO
import json

import pandas as pd
import pytest
from openpyxl import Workbook

from segmentation.errors import AnalysisValidationError
from segmentation.ingestion import load_transactions, normalize_transactions


def csv_source(rows="0007,I1,2020-01-01,2,3\nA9,I2,2020-01-02,1,4\nNA,I3,2020-01-03,1,5\n"):
    return StringIO("CustomerID,InvoiceNo,InvoiceDate,Quantity,UnitPrice\n" + rows)


def test_csv_preserves_identifiers_and_types():
    result = load_transactions(csv_source(), file_format="csv")
    assert result.CustomerID.tolist() == ["0007", "A9", "NA"]
    assert result.InvoiceNo.tolist() == ["I1", "I2", "I3"]
    assert pd.api.types.is_datetime64_any_dtype(result.InvoiceDate)
    assert result.Quantity.tolist() == [2, 1, 1]


def test_csv_path(tmp_path):
    path = tmp_path / "purchases.csv"
    path.write_text(csv_source().getvalue(), encoding="utf-8")
    assert load_transactions(path).CustomerID.iloc[0] == "0007"


def test_xlsx_preserves_stored_text_and_numeric_identifiers():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CustomerID", "InvoiceNo", "InvoiceDate", "Quantity", "UnitPrice"])
    for identifier in ("0007", "A9", 123, "NA"):
        sheet.append([identifier, "0012", pd.Timestamp("2020-01-01").to_pydatetime(), 2, 3])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    result = load_transactions(stream, file_format="xlsx")
    assert result.CustomerID.tolist() == ["0007", "A9", "123", "NA"]
    assert result.InvoiceNo.tolist() == ["0012"] * 4
    assert result.InvoiceDate.iloc[0] == pd.Timestamp("2020-01-01")


def test_xlsx_duplicate_headers_rejected():
    book = Workbook()
    book.active.append(["CustomerID", "CustomerID", "InvoiceDate", "Quantity", "UnitPrice"])
    book.active.append(["A", "I", "2020-01-01", 1, 2])
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(stream, file_format="xlsx")
    assert caught.value.issues[0].code == "duplicate_columns"


def test_missing_columns_structured():
    with pytest.raises(AnalysisValidationError) as caught:
        normalize_transactions(pd.DataFrame({"CustomerID": ["A"]}))
    error = caught.value.to_dict()
    assert error["stage"] == "ingestion"
    assert {i["column"] for i in error["issues"]} == {"InvoiceNo", "InvoiceDate", "Quantity", "UnitPrice"}
    json.dumps(error)


@pytest.mark.parametrize("stream", [StringIO(""), csv_source("")])
def test_empty_csv(stream):
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(stream, file_format="csv")
    assert caught.value.issues[0].code == "empty_input"


def test_empty_xlsx():
    book = Workbook()
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(stream, file_format="xlsx")
    assert caught.value.issues[0].code == "empty_input"


def test_unsupported_format():
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions("input.xls")
    assert caught.value.issues[0].code == "unsupported_format"


def test_duplicate_csv_header():
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(StringIO("CustomerID,CustomerID\na,b\n"), file_format="csv")
    assert caught.value.issues[0].code == "duplicate_columns"


@pytest.mark.parametrize("value", ["01/02/2020", "2020-02-30", "garbage", "", 43831])
def test_invalid_ambiguous_dates(value, transactions):
    transactions["InvoiceDate"] = transactions.InvoiceDate.astype(object)
    transactions.loc[0, "InvoiceDate"] = value
    with pytest.raises(AnalysisValidationError) as caught:
        normalize_transactions(transactions)
    issue = next(i for i in caught.value.issues if i.code == "invalid_date")
    assert issue.column == "InvoiceDate"
    assert issue.rows == (2,)
    json.dumps(caught.value.to_dict())


def test_explicit_date_format():
    result = load_transactions(csv_source("A,I,01/02/2020,1,2\n"), file_format="csv", date_format="%d/%m/%Y")
    assert result.InvoiceDate.iloc[0] == pd.Timestamp("2020-02-01")


def test_timezone_normalization():
    result = load_transactions(csv_source("A,I,2020-01-01T12:00:00+02:00,1,2\n"), file_format="csv")
    assert result.InvoiceDate.iloc[0] == pd.Timestamp("2020-01-01T10:00:00")


@pytest.mark.parametrize("column,value", [("Quantity", "bad"), ("UnitPrice", "bad"), ("Quantity", float("inf")), ("UnitPrice", float("nan")), ("Quantity", True)])
def test_invalid_numeric(column, value, transactions):
    transactions[column] = transactions[column].astype(object)
    transactions.loc[0, column] = value
    with pytest.raises(AnalysisValidationError) as caught:
        normalize_transactions(transactions)
    issue = next(i for i in caught.value.issues if i.code == "invalid_numeric")
    assert issue.column == column
    assert issue.rows == (2,)


def test_missing_customer_is_retained_for_cleaning():
    result = load_transactions(csv_source(",I,2020-01-01,1,2\n"), file_format="csv")
    assert pd.isna(result.CustomerID.iloc[0])


def test_blank_invoice_rejected():
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(csv_source("A,,2020-01-01,1,2\n"), file_format="csv")
    assert any(i.code == "missing_invoice" for i in caught.value.issues)


def test_no_input_mutation(transactions):
    original = transactions.copy(deep=True)
    normalize_transactions(transactions)
    pd.testing.assert_frame_equal(transactions, original)


def test_numeric_ids_normalized_without_changing_text(transactions):
    transactions["CustomerID"] = transactions.CustomerID.astype(object)
    transactions.loc[0, "CustomerID"] = 123.0
    transactions.loc[1, "CustomerID"] = "00123"
    result = normalize_transactions(transactions)
    assert result.CustomerID.iloc[:2].tolist() == ["123", "00123"]


def test_fractional_quantities_are_supported():
    result = load_transactions(csv_source("A,I,2020-01-01,1.5,2\n"), file_format="csv")
    assert result.Quantity.iloc[0] == 1.5


@pytest.mark.parametrize("rows", ["A,I,2020-01-01,1,2,unexpected\n", "A,I,2020-01-01,1\n"])
def test_ragged_csv_rejected(rows):
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(csv_source(rows), file_format="csv")
    assert caught.value.issues[0].code == "invalid_csv_row"


def test_binary_csv_stream():
    stream = BytesIO(csv_source().getvalue().encode("utf-8"))
    assert load_transactions(stream, file_format="csv").CustomerID.iloc[0] == "0007"


def test_corrupt_workbook_is_structured():
    with pytest.raises(AnalysisValidationError) as caught:
        load_transactions(BytesIO(b"not a workbook"), file_format="xlsx")
    assert caught.value.issues[0].code == "read_error"
